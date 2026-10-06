import os
import io
import pytest
from fastapi.testclient import TestClient
import pypdf
import docx

from src.main import app
from database.session import init_db, get_db
from database.models import CompanyDocument, DocumentChunk, DocumentProcessingJob, DocumentStatusEnum

client = TestClient(app)

def create_sample_pdf_bytes() -> bytes:
    """Helper to generate a valid PDF byte stream with text content."""
    pdf_data = (
        b"%PDF-1.4\n"
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n"
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj\n"
        b"4 0 obj << /Length 130 >> stream\n"
        b"BT /F1 24 Tf 100 700 Td (SECTION 1.0: COMPANY HANDBOOK) Tj ET\n"
        b"BT /F1 12 Tf 100 650 Td (All employees must comply with HR policies.) Tj ET\n"
        b"endstream endobj\n"
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n"
        b"xref\n0 6\n"
        b"0000000000 65535 f \n"
        b"0000000009 00000 n \n"
        b"0000000058 00000 n \n"
        b"0000000115 00000 n \n"
        b"0000000261 00000 n \n"
        b"0000000442 00000 n \n"
        b"trailer << /Size 6 /Root 1 0 R >>\n"
        b"startxref\n523\n"
        b"%%EOF\n"
    )
    return pdf_data

def create_sample_docx_bytes() -> bytes:
    """Helper to generate a valid DOCX byte stream with headings and paragraphs."""
    doc = docx.Document()
    doc.add_heading("SECTION 1.0: CODE OF CONDUCT", level=1)
    doc.add_paragraph("All employees must maintain professional conduct and policy compliance.")
    doc.add_heading("SECTION 2.0: INFORMATION SECURITY", level=1)
    doc.add_paragraph("Passwords must be updated every 90 days and never shared externally.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def test_valid_pdf_upload():
    """Test uploading a valid PDF document with page traceability."""
    pdf_bytes = create_sample_pdf_bytes()
    files = {"file": ("policy_handbook.pdf", pdf_bytes, "application/pdf")}
    data = {
        "document_id": "POL-PDF-001",
        "title": "Company Handbook PDF",
        "category": "company_handbook",
        "effective_date": "2026-01-01",
        "version": 1,
        "department": "HR"
    }

    response = client.post("/api/documents/upload", files=files, data=data)
    assert response.status_code == 201, response.text
    res = response.json()
    assert res["document"]["document_id"] == "POL-PDF-001"
    assert res["document"]["file_format"] == "pdf"
    assert res["chunks_count"] >= 1

def test_valid_docx_upload_and_traceability():
    """Test uploading a valid DOCX document and verifying paragraph/section traceability."""
    docx_bytes = create_sample_docx_bytes()
    files = {"file": ("sop_escalation.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {
        "document_id": "SOP-DOCX-001",
        "title": "Escalation Procedure DOCX",
        "category": "department_sop",
        "effective_date": "2026-01-01",
        "version": 1,
        "department": "Customer Support",
        "description": "SOP document for handling customer escalations."
    }

    response = client.post("/api/documents/upload", files=files, data=data)
    assert response.status_code == 201, response.text
    res = response.json()
    assert res["document"]["document_id"] == "SOP-DOCX-001"
    assert res["document"]["file_format"] == "docx"
    assert res["document"]["description"] == "SOP document for handling customer escalations."

    # Test chunk traceability for DOCX
    chunk_res = client.get("/api/documents/SOP-DOCX-001/chunks")
    assert chunk_res.status_code == 200
    chunks = chunk_res.json()
    assert len(chunks) >= 2
    assert any("Para" in (c["paragraph_reference"] or "") for c in chunks)
    assert any("SECTION 1.0" in (c["heading"] or "") for c in chunks)

def test_invalid_file_type():
    """Test rejecting unsupported file formats like .exe or .py."""
    files = {"file": ("script.exe", b"binary content", "application/octet-stream")}
    data = {
        "document_id": "BAD-FILE-001",
        "title": "Malicious Executable",
        "category": "hr_policy",
        "effective_date": "2026-01-01",
        "version": 1
    }
    response = client.post("/api/documents/upload", files=files, data=data)
    assert response.status_code == 400
    assert "Unsupported file extension" in response.json()["detail"]

def test_empty_document():
    """Test rejecting empty 0-byte documents."""
    files = {"file": ("empty.txt", b"", "text/plain")}
    data = {
        "document_id": "EMPTY-001",
        "title": "Empty Document",
        "category": "hr_policy",
        "effective_date": "2026-01-01",
        "version": 1
    }
    response = client.post("/api/documents/upload", files=files, data=data)
    assert response.status_code == 400
    assert "empty (0 bytes)" in response.json()["detail"]

def test_duplicate_document_content():
    """Test detecting duplicate document upload based on SHA-256 content hash."""
    content = b"# SECTION 1.0\nUnique content for hash check."
    files1 = {"file": ("doc1.md", content, "text/markdown")}
    data1 = {
        "document_id": "DOC-HASH-001",
        "title": "Original Document",
        "category": "hr_policy",
        "effective_date": "2026-01-01",
        "version": 1
    }
    res1 = client.post("/api/documents/upload", files=files1, data=data1)
    assert res1.status_code == 201

    # Upload exact same content with different document ID
    files2 = {"file": ("doc2.md", content, "text/markdown")}
    data2 = {
        "document_id": "DOC-HASH-002",
        "title": "Duplicate Content Document",
        "category": "hr_policy",
        "effective_date": "2026-01-01",
        "version": 1
    }
    res2 = client.post("/api/documents/upload", files=files2, data=data2)
    assert res2.status_code == 409
    assert "Duplicate document content detected" in res2.json()["detail"]

def test_version_handling_and_activation():
    """Test uploading multiple versions of a document, preserving history, and activating versions."""
    content_v1 = b"# SECTION 1.0\nVersion 1 Content."
    files_v1 = {"file": ("doc_v1.md", content_v1, "text/markdown")}
    data_v1 = {
        "document_id": "VER-DOC-001",
        "title": "Versioned Document",
        "category": "hr_policy",
        "effective_date": "2026-01-01",
        "version": 1
    }
    res_v1 = client.post("/api/documents/upload", files=files_v1, data=data_v1)
    assert res_v1.status_code == 201

    content_v2 = b"# SECTION 1.0\nVersion 2 Content Updated."
    files_v2 = {"file": ("doc_v2.md", content_v2, "text/markdown")}
    data_v2 = {
        "document_id": "VER-DOC-001",
        "title": "Versioned Document Updated",
        "category": "hr_policy",
        "effective_date": "2026-02-01",
        "version": 2
    }
    res_v2 = client.post("/api/documents/upload", files=files_v2, data=data_v2)
    assert res_v2.status_code == 201
    assert res_v2.json()["superseded_previous_version"] is True

    # Check version history listing
    ver_res = client.get("/api/documents/VER-DOC-001/versions")
    assert ver_res.status_code == 200
    versions = ver_res.json()
    assert len(versions) == 2
    assert versions[0]["version"] == 2
    assert versions[0]["is_active"] is True
    assert versions[1]["version"] == 1
    assert versions[1]["is_active"] is False

    # Test activating historical version 1
    act_res = client.post("/api/documents/VER-DOC-001/versions/1/activate")
    assert act_res.status_code == 200
    assert act_res.json()["activated_version"] == 1

    # Verify active status
    ver_res_after = client.get("/api/documents/VER-DOC-001/versions")
    versions_after = ver_res_after.json()
    v1_item = next(v for v in versions_after if v["version"] == 1)
    assert v1_item["is_active"] is True

def test_unauthorized_upload():
    """Test rejecting document management operations when user role is 'employee'."""
    files = {"file": ("test.md", b"# Policy", "text/markdown")}
    data = {
        "document_id": "UNAUTH-001",
        "title": "Unauthorized Test",
        "category": "hr_policy",
        "effective_date": "2026-01-01",
        "version": 1
    }
    headers = {"X-User-Role": "employee"}
    response = client.post("/api/documents/upload", files=files, data=data, headers=headers)
    assert response.status_code == 403
    assert "Access Denied" in response.json()["detail"]

def test_reprocess_document():
    """Test reprocessing an existing document version."""
    content = b"# SECTION 1.0: REPROCESS TEST\nTesting document reprocess endpoint."
    files = {"file": ("reprocess.md", content, "text/markdown")}
    data = {
        "document_id": "REPROC-001",
        "title": "Reprocess Test Document",
        "category": "hr_policy",
        "effective_date": "2026-01-01",
        "version": 1
    }
    client.post("/api/documents/upload", files=files, data=data)

    reproc_res = client.post("/api/documents/REPROC-001/versions/1/process")
    assert reproc_res.status_code == 200
    job = reproc_res.json()
    assert job["status"] == "completed"
    assert job["total_chunks"] >= 1
